"""
Tender Intelligence — Web UI

Streamlit dashboard for managing the CGI tender intelligence system.
Run with: streamlit run app.py
"""

import streamlit as st
import json
import os
import subprocess
import sys
from datetime import datetime

# Must be first Streamlit call
st.set_page_config(
    page_title="CGI Tender Intelligence",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Import project modules
from storage import (
    get_db, init_db, get_open_tenders, get_stats, get_competitors,
    get_price_history, get_unanalyzed_results,
)
from routing import load_routing_rules, CONFIG_PATH, create_default_config

init_db()

# ── Sidebar ──────────────────────────────────────────────────────────────

st.sidebar.title("CGI Tender Intelligence")
st.sidebar.markdown("---")

page = st.sidebar.radio("Navigation", [
    "Dashboard",
    "Open Tenders",
    "Competitor Intelligence",
    "Price Intelligence",
    "Routing & Notifications",
    "Settings & Run",
])

st.sidebar.markdown("---")
stats = get_stats()
st.sidebar.metric("Total Tenders", stats["total"])
st.sidebar.metric("Open", stats["open"])
st.sidebar.metric("Awaiting Notification", stats["unnotified"])


# ── Dashboard ────────────────────────────────────────────────────────────

if page == "Dashboard":
    st.title("Dashboard")

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Total Tenders", stats["total"])
    col2.metric("Open Tenders", stats["open"])
    col3.metric("Closed/Results", stats["closed"])
    col4.metric("Awaiting Notification", stats["unnotified"])

    st.markdown("---")

    # Category breakdown
    col_left, col_right = st.columns(2)

    with col_left:
        st.subheader("Open Tenders by Category")
        if stats.get("by_category"):
            for cat, count in sorted(stats["by_category"].items(), key=lambda x: -x[1]):
                label = {
                    "open_competition": "Open Competitions",
                    "dynamic_purchasing": "Dynamic Purchasing Systems",
                    "early_signal": "Early Signals",
                    "direct_award": "Direct Awards",
                    "unknown": "Uncategorized",
                }.get(cat, cat)
                st.markdown(f"**{label}**: {count}")
        else:
            st.info("No data yet. Run the scraper first.")

    with col_right:
        st.subheader("Top Competitors")
        competitors = get_competitors(min_wins=1)
        if competitors:
            for c in competitors[:10]:
                avg = f"€{c['avg_contract_value']:,.0f}" if c["avg_contract_value"] else ""
                st.markdown(f"**{c['name']}** — {c['wins']} wins {avg}")
        else:
            st.info("No competitor data yet. Run analysis on award results.")

    st.markdown("---")

    # Recent tenders
    st.subheader("Recent Open Tenders")
    open_tenders = get_open_tenders()
    if open_tenders:
        for t in open_tenders[:10]:
            with st.expander(f"{t['name'][:80]}"):
                st.markdown(f"**Organisation:** {t.get('organisation', 'N/A')}")
                st.markdown(f"**Type:** {t.get('type', 'N/A')}")
                st.markdown(f"**Deadline:** {t.get('deadline', 'N/A')}")
                st.markdown(f"**Category:** {t.get('category', 'N/A')}")
                if t.get("url"):
                    st.markdown(f"[Open on tarjouspalvelu.fi]({t['url']})")
                desc = t.get("description", "")
                if desc:
                    st.markdown(f"**Description:** {desc[:500]}...")
    else:
        st.info("No tenders yet. Run the scraper first.")


# ── Open Tenders ─────────────────────────────────────────────────────────

elif page == "Open Tenders":
    st.title("Open Tenders")

    open_tenders = get_open_tenders()

    if not open_tenders:
        st.info("No open tenders. Run the scraper first.")
    else:
        # Filters
        col1, col2, col3 = st.columns(3)
        categories = sorted(set(t.get("category", "") for t in open_tenders))
        orgs = sorted(set(t.get("organisation", "") for t in open_tenders if t.get("organisation")))

        with col1:
            cat_filter = st.selectbox("Category", ["All"] + categories)
        with col2:
            org_filter = st.selectbox("Organisation", ["All"] + orgs)
        with col3:
            search = st.text_input("Search", placeholder="Search tender name or description...")

        # Apply filters
        filtered = open_tenders
        if cat_filter != "All":
            filtered = [t for t in filtered if t.get("category") == cat_filter]
        if org_filter != "All":
            filtered = [t for t in filtered if t.get("organisation") == org_filter]
        if search:
            search_lower = search.lower()
            filtered = [t for t in filtered if
                        search_lower in t.get("name", "").lower() or
                        search_lower in t.get("description", "").lower()]

        st.markdown(f"**Showing {len(filtered)} of {len(open_tenders)} tenders**")

        for t in filtered:
            with st.expander(f"{'🟢' if t.get('category') == 'open_competition' else '🟡'} {t['name'][:90]}"):
                col_a, col_b = st.columns([2, 1])
                with col_a:
                    st.markdown(f"**Organisation:** {t.get('organisation', 'N/A')}")
                    st.markdown(f"**Type:** {t.get('type', 'N/A')}")
                    desc = t.get("description", "")
                    if desc:
                        st.markdown(f"{desc[:600]}")
                with col_b:
                    st.markdown(f"**Category:** `{t.get('category', 'N/A')}`")
                    st.markdown(f"**Published:** {t.get('published', 'N/A')}")
                    st.markdown(f"**Deadline:** {t.get('deadline', 'N/A')}")
                    if t.get("url"):
                        st.link_button("Open on tarjouspalvelu.fi", t["url"])


# ── Competitor Intelligence ──────────────────────────────────────────────

elif page == "Competitor Intelligence":
    st.title("Competitor Intelligence")

    competitors = get_competitors(min_wins=0)

    if not competitors:
        st.info("No competitor data yet. Run the scraper and analysis first.")
    else:
        st.markdown(f"**{len(competitors)} companies tracked**")

        # Summary table
        table_data = []
        for c in competitors:
            table_data.append({
                "Company": c["name"],
                "Wins": c["wins"],
                "Total Value": f"€{c['total_contract_value']:,.0f}" if c["total_contract_value"] else "N/A",
                "Avg Value": f"€{c['avg_contract_value']:,.0f}" if c["avg_contract_value"] else "N/A",
                "Sectors": c.get("sectors", ""),
                "Last Win": c.get("last_win_date", "N/A"),
            })

        st.dataframe(table_data, use_container_width=True, hide_index=True)


# ── Price Intelligence ───────────────────────────────────────────────────

elif page == "Price Intelligence":
    st.title("Price Intelligence")

    prices = get_price_history()

    if not prices:
        st.info("No price data yet. Run the scraper and analysis first.")
    else:
        # Group by sector
        by_sector: dict[str, list] = {}
        for p in prices:
            sector = p.get("sector", "Other")
            value = p.get("winning_price") or p.get("estimated_value") or 0
            if value > 0:
                if sector not in by_sector:
                    by_sector[sector] = []
                by_sector[sector].append({
                    "tender": p.get("tender_name", ""),
                    "winner": p.get("winner_name", "N/A"),
                    "value": value,
                })

        for sector, items in sorted(by_sector.items(), key=lambda x: -len(x[1])):
            values = [i["value"] for i in items]
            avg = sum(values) / len(values)

            st.subheader(f"{sector}")
            col1, col2, col3 = st.columns(3)
            col1.metric("Contracts", len(items))
            col2.metric("Avg Value", f"€{avg:,.0f}")
            col3.metric("Range", f"€{min(values):,.0f} — €{max(values):,.0f}")

            with st.expander(f"Details ({len(items)} contracts)"):
                for item in items:
                    st.markdown(f"- **{item['tender'][:60]}** — €{item['value']:,.0f} → {item['winner']}")

            st.markdown("---")


# ── Routing & Notifications ──────────────────────────────────────────────

elif page == "Routing & Notifications":
    st.title("Routing & Notifications")

    st.subheader("Current Routing Rules")
    st.markdown(f"Config file: `{CONFIG_PATH}`")

    rules = load_routing_rules()
    if rules:
        for i, rule in enumerate(rules):
            with st.expander(
                f"{'🌐 Catch-all' if rule['is_catchall'] else rule['department']}"
                f" → {rule['email']}"
                f" ({len(rule['keywords'])} keywords)"
            ):
                col1, col2 = st.columns(2)
                with col1:
                    st.markdown(f"**Department:** {rule['department']}")
                    st.markdown(f"**Contact:** {rule['contact']}")
                    st.markdown(f"**Email:** {rule['email']}")
                with col2:
                    st.markdown(f"**Keywords:** {', '.join(rule['keywords'])}")
                    st.markdown(f"**Notes:** {rule.get('notes', '')}")
    else:
        st.warning("No routing rules found.")

    st.markdown("---")

    st.subheader("Edit Routing Rules")

    # Editable routing table
    if "routing_rules" not in st.session_state:
        st.session_state.routing_rules = [
            {
                "keywords": ", ".join(r["keywords"]),
                "department": r["department"],
                "contact": r["contact"],
                "email": r["email"],
            }
            for r in rules
        ]

    edited = st.data_editor(
        st.session_state.routing_rules,
        num_rows="dynamic",
        use_container_width=True,
        column_config={
            "keywords": st.column_config.TextColumn("Keywords (comma-separated)", width="large"),
            "department": st.column_config.TextColumn("Department"),
            "contact": st.column_config.TextColumn("Contact"),
            "email": st.column_config.TextColumn("Email"),
        },
    )

    if st.button("Save Routing Rules", type="primary"):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Routing Rules"
        ws["A1"] = "Keywords"
        ws["B1"] = "Department"
        ws["C1"] = "Contact"
        ws["D1"] = "Email"
        ws["E1"] = "Notes"
        for i, row in enumerate(edited, start=2):
            ws[f"A{i}"] = row.get("keywords", "")
            ws[f"B{i}"] = row.get("department", "")
            ws[f"C{i}"] = row.get("contact", "")
            ws[f"D{i}"] = row.get("email", "")
        wb.save(CONFIG_PATH)
        st.success(f"Saved {len(edited)} rules to {CONFIG_PATH}")
        st.session_state.routing_rules = edited


# ── Settings & Run ───────────────────────────────────────────────────────

elif page == "Settings & Run":
    st.title("Settings & Run")

    # Settings
    st.subheader("Pipeline Settings")

    col1, col2 = st.columns(2)
    with col1:
        org = st.selectbox("Organization to scrape", [
            "helsinki", "espoo", "vantaa", "tampere", "oulu", "turku"
        ])
        enable_summarization = st.toggle("AI Summarization", value=True)
        summarize_count = st.slider("Max tenders to summarize", 5, 50, 10) if enable_summarization else 0

    with col2:
        enable_email = st.toggle("Send Email Notifications", value=True)
        enable_analysis = st.toggle("Award Results Analysis", value=True)
        analysis_count = st.slider("Max results to analyze", 5, 50, 20) if enable_analysis else 0

    st.markdown("---")

    # Environment status
    st.subheader("Environment")
    env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
    from dotenv import load_dotenv
    load_dotenv(env_path)

    col1, col2, col3 = st.columns(3)
    with col1:
        has_openai = bool(os.getenv("OPENAI_API_KEY"))
        st.markdown(f"**OpenAI API:** {'✅ Configured' if has_openai else '❌ Not set'}")
    with col2:
        has_smtp = bool(os.getenv("SMTP_USER") and os.getenv("SMTP_PASSWORD"))
        st.markdown(f"**SMTP Email:** {'✅ Configured' if has_smtp else '❌ Not set'}")
    with col3:
        has_login = bool(os.getenv("TARJOUSPALVELU_EMAIL")) and not os.getenv("TARJOUSPALVELU_EMAIL", "").startswith("FILL")
        st.markdown(f"**Tarjouspalvelu Login:** {'✅ Configured' if has_login else '⏳ Pending'}")

    st.markdown("---")

    # Run scraper
    st.subheader("Run Scraper")

    if st.button("Run Full Pipeline", type="primary", use_container_width=True):
        with st.status("Running pipeline...", expanded=True) as status:
            st.write(f"Scraping {org}...")

            # Build command
            cmd = [sys.executable, "demo_scraper.py", org]
            env = os.environ.copy()
            env["SUMMARIZE_COUNT"] = str(summarize_count)
            env["ANALYSIS_COUNT"] = str(analysis_count)
            env["ENABLE_EMAIL"] = "1" if enable_email else "0"
            env["ENABLE_SUMMARIZATION"] = "1" if enable_summarization else "0"
            env["ENABLE_ANALYSIS"] = "1" if enable_analysis else "0"

            # Run and stream output
            process = subprocess.Popen(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                cwd=os.path.dirname(os.path.abspath(__file__)),
                env=env,
            )

            output_container = st.empty()
            full_output = ""
            for line in process.stdout:
                full_output += line
                output_container.code(full_output[-3000:], language="text")

            process.wait()

            if process.returncode == 0:
                status.update(label="Pipeline complete!", state="complete")
            else:
                status.update(label="Pipeline failed", state="error")

    st.markdown("---")

    # Database management
    st.subheader("Database")
    db_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tenders.db")
    if os.path.exists(db_path):
        size = os.path.getsize(db_path) / 1024
        st.markdown(f"**Database:** `tenders.db` ({size:.0f} KB)")

        if st.button("Reset Database", type="secondary"):
            os.remove(db_path)
            init_db()
            st.success("Database reset.")
            st.rerun()
    else:
        st.info("No database yet. Run the scraper to create it.")
